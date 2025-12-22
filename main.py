import os
import re
import threading
import tempfile
import tkinter as tk
from tkinter import filedialog, messagebox
import ttkbootstrap as tb 
from ttkbootstrap.constants import *
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx2python import docx2python

# ----------------------------
# Utilities
# ----------------------------

def parse_number_input(text):
    numbers = set()
    for part in text.split(","):
        part = part.strip()
        if not part: continue
        if "-" in part:
            try:
                start, end = part.split("-")
                numbers.update(range(int(start), int(end) + 1))
            except ValueError: continue
        else:
            try: numbers.add(int(part))
            except ValueError: continue
    return sorted(numbers)

def normalize_answer_text(text):
    match = re.match(r'^([A-Ea-e])[\.\)]\s*(.*)', text)
    if match:
        letter = match.group(1).upper()
        content = match.group(2).strip()
        return f"{letter}. {content}"
    return text.strip()

# ----------------------------
# Document Processing
# ----------------------------

def process_document(file_path):
    with docx2python(file_path) as doc:
        text = doc.text
        images = doc.images 

    image_lookup = {os.path.basename(k): v for k, v in images.items()}

    questions = {}
    current_q = None
    
    # EXACT REGEX FROM YOUR WORKING VERSION
    q_pattern = re.compile(r'^(\d+)\.?\s*(?:\([^\)]*\))?\s*(.*)')
    ans_pattern = re.compile(r'^([A-Ea-e])[\.\)]')

    for line in text.splitlines():
        line = line.strip()
        if not line: continue

        q_match = q_pattern.match(line)
        if q_match:
            current_q = int(q_match.group(1))
            questions[current_q] = {
                "header_parts": [q_match.group(2).strip()],
                "elements": [],
                "found_first_answer": False
            }
            continue

        if current_q is not None:
            img_match = re.search(r'----(.*?image\d+\.\w+)----', line)
            if img_match:
                img_filename = os.path.basename(img_match.group(1))
                questions[current_q]["elements"].append({
                    "type": "image",
                    "name": img_filename
                })
                continue

            trash_patterns = ["Image alt text", "AI-generated content", "----image"]
            if any(p in line for p in trash_patterns):
                continue

            if ans_pattern.match(line):
                questions[current_q]["found_first_answer"] = True
                questions[current_q]["elements"].append({
                    "type": "text",
                    "content": line
                })
            elif not questions[current_q]["found_first_answer"]:
                questions[current_q]["header_parts"].append(line)
            else:
                questions[current_q]["elements"].append({
                    "type": "text",
                    "content": line
                })

    for q_id in questions:
        clean_parts = [p.strip() for p in questions[q_id]["header_parts"] if p.strip() and p != "?"]
        questions[q_id]["header"] = " ".join(clean_parts) if clean_parts else "Question Text Missing"

    return questions, image_lookup

# ----------------------------
# UI Class
# ----------------------------

class SynapseSorterApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Synapse Sorter")
        self.root.geometry("550x650")
        self.questions, self.image_dict = None, None

        try:
            self.root.iconbitmap("icon.ico")
        except:
            pass
            
        self.questions, self.image_dict = None, None
        
        self.frame = tb.Frame(self.root, padding=30)
        self.frame.pack(fill="both", expand=True)
        
        tb.Label(self.frame, text="SYNAPSE SORTER", font=("Helvetica", 20, "bold"), bootstyle="primary").pack(pady=(0, 25))
        
        tb.Label(self.frame, text="Exam Title:", font=("Helvetica", 10, "bold")).pack(anchor="w")
        self.title_ent = tb.Entry(self.frame, bootstyle="primary")
        self.title_ent.insert(0, "Immunology Assessment")
        self.title_ent.pack(fill="x", pady=(5, 20))
        
        self.load_btn = tb.Button(self.frame, text="LOAD QUESTION BANK", command=self.load_file, bootstyle="outline-primary")
        self.load_btn.pack(fill="x", pady=5)
        self.status = tb.Label(self.frame, text="No file selected", font=("Helvetica", 9), bootstyle="secondary")
        self.status.pack(pady=(0, 20))
        
        tb.Label(self.frame, text="Select Question Numbers:", font=("Helvetica", 10, "bold")).pack(anchor="w")
        tb.Label(self.frame, text="Example: 1-5, 8, 20-26", font=("Helvetica", 8), bootstyle="secondary").pack(anchor="w")
        self.range_ent = tb.Entry(self.frame, bootstyle="primary")
        self.range_ent.pack(fill="x", pady=5)
        
        self.progress = tb.Progressbar(self.frame, bootstyle="primary-striped", value=0)
        self.progress.pack(fill="x", pady=20)
        
        self.gen_btn = tb.Button(self.frame, text="GENERATE CLEAN EXAM", command=self.generate, bootstyle="primary")
        self.gen_btn.pack(fill="x", pady=10)
        
        tb.Button(self.frame, text="HELP / INSTRUCTIONS", command=self.show_help, bootstyle="link").pack(pady=5)
        tb.Label(self.frame, text="Developed for Tamara", font=("Helvetica", 8, "italic"), bootstyle="secondary").pack(side="bottom", pady=10)

    def show_help(self):
        help_msg = (
            "1. Enter an Exam Title.\n"
            "2. Click LOAD to select your source Word document.\n"
            "3. Enter the question numbers you want (e.g., 1-10, 15).\n"
            "4. Click GENERATE to save your new formatted exam.\n\n"
            "Bilingual text and images are restored automatically."
        )
        messagebox.showinfo("Synapse Sorter Help", help_msg)

    def load_file(self):
        path = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
        if path:
            self.questions, self.image_dict = process_document(path)
            self.status.config(text=f"Loaded {len(self.questions)} questions", bootstyle="success")

    def generate(self):
        if not self.questions: return
        nums = parse_number_input(self.range_ent.get())
        save_path = filedialog.asksaveasfilename(defaultextension=".docx")
        if save_path:
            self.progress["maximum"] = len(nums)
            self.progress["value"] = 0
            threading.Thread(target=self.run_task, args=(save_path, self.title_ent.get(), nums), daemon=True).start()

    def run_task(self, path, title, nums):
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(11)
        
        title_p = doc.add_heading(title, level=0)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        with tempfile.TemporaryDirectory() as tmp:
            new_num = 1
            for n in nums:
                if n in self.questions:
                    q = self.questions[n]
                    header_p = doc.add_paragraph()
                    header_p.paragraph_format.space_before = Pt(12)
                    header_p.add_run(f"{new_num}. {q['header']}").bold = True
                    for e in q["elements"]:
                        if e["type"] == "text":
                            ans = doc.add_paragraph(normalize_answer_text(e["content"]))
                            ans.paragraph_format.left_indent = Inches(0.3)
                        elif e["type"] == "image":
                            img_name = e["name"]
                            if img_name in self.image_dict:
                                img_path = os.path.join(tmp, img_name)
                                with open(img_path, "wb") as f: f.write(self.image_dict[img_name])
                                img_p = doc.add_paragraph()
                                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                img_p.add_run().add_picture(img_path, width=Inches(2.5))
                    new_num += 1
                self.progress.step(1)
        doc.save(path)
        messagebox.showinfo("Success", "Exam Generated!")

if __name__ == "__main__":
    root_win = tb.Window(themename="flatly")
    app = SynapseSorterApp(root_win)
    root_win.mainloop()