import copy
import re

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# ----------------------------
# Patterns
# ----------------------------

# A question starts with its number, e.g. "12. ..." or "12) ..."
QUESTION_RE = re.compile(r'^(\s*)(\d+)(\s*)([\.\)])')
# Lines starting with this mark are section titles and are never copied
SECTION_MARK = "*"

R_ATTRS = (qn('r:embed'), qn('r:link'), qn('r:id'))


# ----------------------------
# Number Input
# ----------------------------

def parse_number_input(text):
    """Turn '20-26, 3, 8' into [20, ..., 26, 3, 8], keeping the order typed."""
    numbers = []
    for part in text.split(","):
        part = part.strip()
        if not part: continue
        if "-" in part:
            try:
                start, end = (int(x) for x in part.split("-"))
            except ValueError: continue
            step = 1 if end >= start else -1
            numbers.extend(range(start, end + step, step))
        else:
            try: numbers.append(int(part))
            except ValueError: continue
    return numbers


def _cell_to_number(value, coordinate):
    if isinstance(value, bool):
        raise ValueError(f"Cell {coordinate}: '{value}' is not a question number")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    raise ValueError(f"Cell {coordinate}: '{value}' is not a question number")


def read_permutations(xlsx_path):
    """Read an Excel sheet where each column is one exam version (permuta).

    Row 1 holds the version names (e.g. 'Permuta_1'), the rows below hold the
    question numbers in the order they should appear. Returns a list of
    (name, [numbers]) tuples, one per non-empty column.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows())
    finally:
        wb.close()
    if not rows:
        return []

    header = rows[0]
    has_header = any(isinstance(c.value, str) and not c.value.strip().isdigit() for c in header)
    data_rows = rows[1:] if has_header else rows

    permutations = []
    for col in range(max(len(r) for r in rows)):
        name = None
        if has_header and col < len(header) and header[col].value is not None:
            name = str(header[col].value).strip()
        numbers = []
        for row in data_rows:
            if col >= len(row) or row[col].value is None: continue
            if isinstance(row[col].value, str) and not row[col].value.strip(): continue
            numbers.append(_cell_to_number(row[col].value, row[col].coordinate))
        if numbers:
            permutations.append((name or f"Permuta_{len(permutations) + 1}", numbers))
    return permutations


# ----------------------------
# Document Processing
# ----------------------------

def _text(el):
    return "".join(t.text or "" for t in el.iter(qn('w:t')))


def _has_list_numbering(p):
    ppr = p.find(qn('w:pPr'))
    return ppr is not None and ppr.find(qn('w:numPr')) is not None


def _is_blank(el):
    if _text(el).strip():
        return False
    return next(el.iter(qn('w:drawing'), qn('w:pict'), qn('w:object')), None) is None


class QuestionBank:
    """Where each numbered question lives inside the source .docx.

    Questions are stored as positions of body elements, so every exam can be
    built from a fresh copy of the source and keep its styles, list
    numbering and images.
    """

    def __init__(self, path):
        self.path = path
        self.questions = {}     # number -> [body element index, ...]
        self.duplicates = []    # numbers that appear more than once in the bank

        body = Document(path).element.body
        current = None
        for i, el in enumerate(body.iterchildren()):
            if el.tag not in (qn('w:p'), qn('w:tbl')): continue
            text = _text(el)

            if el.tag == qn('w:p') and text.strip().startswith(SECTION_MARK):
                current = None
                continue

            match = QUESTION_RE.match(text) if el.tag == qn('w:p') and not _has_list_numbering(el) else None
            if match:
                current = int(match.group(2))
                if current in self.questions:
                    self.duplicates.append(current)
                self.questions[current] = [i]
                continue

            if current is not None:
                self.questions[current].append(i)

        # Trailing blank paragraphs are spacing, not part of the question
        children = list(body.iterchildren())
        for indexes in self.questions.values():
            while len(indexes) > 1 and _is_blank(children[indexes[-1]]):
                indexes.pop()

    def __len__(self):
        return len(self.questions)

    def is_empty(self, number, children):
        indexes = self.questions[number]
        stem = QUESTION_RE.sub("", _text(children[indexes[0]]), count=1)
        return not stem.strip() and all(_is_blank(children[i]) for i in indexes[1:])


def _renumber(p, new_number):
    """Replace the leading question number, which Word often splits over several runs."""
    t_nodes = list(p.iter(qn('w:t')))
    match = QUESTION_RE.match("".join(t.text or "" for t in t_nodes))
    if not match or not t_nodes: return
    to_remove = match.end()
    for t in t_nodes:
        if to_remove <= 0: break
        text = t.text or ""
        cut = min(len(text), to_remove)
        t.text = text[cut:]
        # Without this Word drops the space that now starts the text ("1.Adaptive")
        t.set(qn('xml:space'), 'preserve')
        to_remove -= cut
    t_nodes[0].text = f"{new_number}{match.group(4)}{t_nodes[0].text}"
    t_nodes[0].set(qn('xml:space'), 'preserve')


def _restart_lists(elements, numbering):
    """Give each copied question its own auto-numbered list so answers restart at a."""
    remap = {}
    for num_id in (n for el in elements for n in el.iter(qn('w:numId'))):
        old = num_id.get(qn('w:val'))
        if old == "0": continue
        if old not in remap:
            try:
                source = numbering.num_having_numId(int(old))
            except KeyError:
                continue
            new_num = numbering.add_num(source.abstractNumId.val)
            for ilvl in range(9):
                new_num.add_lvlOverride(ilvl=ilvl).add_startOverride(_level_start(numbering, source.abstractNumId.val, ilvl))
            remap[old] = str(new_num.numId)
        num_id.set(qn('w:val'), remap[old])


def _level_start(numbering, abstract_id, ilvl):
    found = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{ilvl}"]/w:start/@w:val')
    return int(found[0]) if found else 1


def _is_red(hex_value):
    if not hex_value or len(hex_value) != 6: return False
    try:
        r, g, b = (int(hex_value[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    return r >= 0xC0 and g < 0x50 and b < 0x50


def _remove_red(elements):
    """Strip the red colour used to mark correct answers (for student copies)."""
    for el in elements:
        for color in list(el.iter(qn('w:color'))):
            if _is_red(color.get(qn('w:val'))):
                color.getparent().remove(color)


def _drop_unused_images(doc):
    used = {v for el in doc.element.body.iter() for k, v in el.attrib.items() if k in R_ATTRS}
    for r_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.IMAGE and r_id not in used:
            doc.part.rels.pop(r_id)


def _add_title(doc, title):
    try:
        p = doc.add_heading(title, level=0)
    except KeyError:  # source has no Title style
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_exam(bank, numbers, title, out_path, remove_answer_marks=False, on_progress=None):
    """Write an exam with the given questions, in the given order, renumbered from 1.

    Returns a report dict: 'missing' (not in the bank), 'repeated' (asked for
    more than once; only the first is used) and 'empty' (no text in the bank).
    """
    report = {"missing": [], "repeated": [], "empty": []}
    doc = Document(bank.path)
    body = doc.element.body
    children = list(body.iterchildren())
    numbering = doc.part.numbering_part.element

    sect_pr = body.find(qn('w:sectPr'))
    for el in children:
        if el is not sect_pr:
            body.remove(el)
    _add_title(doc, title)

    seen = set()
    new_number = 1
    for step, n in enumerate(numbers, start=1):
        if n in seen:
            report["repeated"].append(n)
        elif n not in bank.questions:
            report["missing"].append(n)
        else:
            seen.add(n)
            if bank.is_empty(n, children):
                report["empty"].append(n)
            copies = [copy.deepcopy(children[i]) for i in bank.questions[n]]
            _renumber(copies[0], new_number)
            _restart_lists(copies, numbering)
            if remove_answer_marks:
                _remove_red(copies)
            for el in copies + [OxmlElement('w:p')]:
                if sect_pr is not None: sect_pr.addprevious(el)
                else: body.append(el)
            new_number += 1
        if on_progress: on_progress(step)

    _drop_unused_images(doc)
    doc.save(out_path)
    return report


def safe_filename(text):
    return re.sub(r'[<>:"/\\|?*]+', "_", text).strip(" .") or "Exam"


def describe_report(report):
    lines = []
    if report["missing"]:
        lines.append("Not found in the question bank: " + ", ".join(map(str, report["missing"])))
    if report["repeated"]:
        lines.append("Listed more than once (used only once): " + ", ".join(map(str, report["repeated"])))
    if report["empty"]:
        lines.append("Empty in the question bank: " + ", ".join(map(str, report["empty"])))
    return "\n".join(lines)
