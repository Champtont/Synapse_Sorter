import struct
import sys
import zlib
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


def tiny_png(color):
    """A valid 2x2 PNG, so tests don't need image files on disk."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + bytes(color) * 2 for _ in range(2))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def add_lettered_list(doc):
    """Add an 'a.' auto-numbered list definition, like the ones Word creates. Returns its numId."""
    numbering = doc.part.numbering_part.element
    abstract_id = 900 + len(numbering.findall(qn('w:abstractNum')))
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/>'
        '<w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'))
    return numbering.add_num(abstract_id).numId


def add_answer(doc, text, num_id, correct=False):
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'))
    run = p.add_run(text)
    if correct:
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    return p


@pytest.fixture
def bank_path(tmp_path):
    """A small question bank shaped like the real one.

    - '*' section titles before and between questions
    - questions 1-3 with auto-lettered answers, correct answer in red
    - question 2 with italic text and an image
    - questions 3 and 4 sharing one list (as Q98 and Q302 do in the real pool)
    - question 5 with answers typed into the same paragraph ('a. ...')
    - question 6 with no text (as Q199 in the real pool)
    - unused image after the last question is never selected
    """
    doc = Document()
    doc.add_paragraph().add_run("*SECTION ONE").bold = True

    def question(number, text, italic=False, split_number=False):
        p = doc.add_paragraph()
        if split_number:
            # As Word saves it: "2" then ". Second..." with no xml:space="preserve"
            p.add_run(str(number))
            p.add_run(". " + text).italic = italic
        else:
            p.add_run(f"{number}. {text}").italic = italic
        return p

    q1_list = add_lettered_list(doc)
    question(1, "First question / Primera pregunta")
    add_answer(doc, "Wrong one", q1_list)
    add_answer(doc, "Right one", q1_list, correct=True)
    doc.add_paragraph()

    q2_list = add_lettered_list(doc)
    question(2, "Second question with picture", italic=True, split_number=True)
    doc.add_picture(_png_file(tmp_path, "used.png", (0, 128, 0)), width=Inches(1))
    add_answer(doc, "Second right", q2_list, correct=True)
    add_answer(doc, "Second wrong", q2_list)
    doc.add_paragraph()
    doc.add_paragraph().add_run("*SECTION TWO").bold = True

    shared = add_lettered_list(doc)
    question(3, "Third question")
    add_answer(doc, "Third A", shared)
    add_answer(doc, "Third B", shared, correct=True)
    doc.add_paragraph()
    question(4, "Fourth question")
    add_answer(doc, "Fourth A", shared, correct=True)
    add_answer(doc, "Fourth B", shared)
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.add_run("5) Fifth question")
    p.add_run().add_break()
    p.add_run("a. Fifth answer A")
    doc.add_paragraph("e. Fifth answer E")
    doc.add_paragraph()

    doc.add_paragraph("6. ")
    doc.add_paragraph()
    doc.add_paragraph().add_run("*SEMINARS").bold = True
    doc.add_picture(_png_file(tmp_path, "unused.png", (0, 0, 255)), width=Inches(1))

    path = tmp_path / "bank.docx"
    doc.save(path)
    return path


def _png_file(tmp_path, name, color):
    path = tmp_path / name
    path.write_bytes(tiny_png(color))
    return str(path)


@pytest.fixture
def make_xlsx(tmp_path):
    def make(rows, name="perms.xlsx"):
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        path = tmp_path / name
        wb.save(path)
        return path
    return make
