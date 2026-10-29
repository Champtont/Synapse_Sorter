import os
import re

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from synapse_core import (QuestionBank, build_exam, describe_report, parse_number_input,
                          read_permutations, safe_filename)


# ----------------------------
# Helpers
# ----------------------------

def build(bank_path, tmp_path, numbers, **kwargs):
    bank = QuestionBank(str(bank_path))
    out = tmp_path / "exam.docx"
    report = build_exam(bank, numbers, "Test Exam", str(out), **kwargs)
    return Document(str(out)), report


def texts(doc):
    return [p.text for p in doc.paragraphs if p.text.strip()]


def question_lines(doc):
    return [t for t in texts(doc) if re.match(r"^\d+[\.\)]", t)]


def num_ids(doc, text):
    return [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.text.startswith(text)]


def red_runs(doc):
    return [r.text for p in doc.paragraphs for r in p.runs
            if r.font.color and r.font.color.type and str(r.font.color.rgb) == "FF0000"]


# ----------------------------
# Number input
# ----------------------------

@pytest.mark.parametrize("text, expected", [
    ("1-3, 8", [1, 2, 3, 8]),
    ("20, 3, 8", [20, 3, 8]),             # order typed is kept
    ("5-3", [5, 4, 3]),                    # reversed range counts down
    ("2, x, 4-, , 7", [2, 7]),             # junk is skipped
    ("3, 3", [3, 3]),                      # repeats are reported by build_exam
    ("", []),
])
def test_parse_number_input(text, expected):
    assert parse_number_input(text) == expected


# ----------------------------
# Excel permutations
# ----------------------------

def test_read_permutations_keeps_columns_and_order(make_xlsx):
    path = make_xlsx([["Permuta_1", "Permuta_2"], [297, 106], [15, 108.0], ["23", None]])
    assert read_permutations(path) == [("Permuta_1", [297, 15, 23]), ("Permuta_2", [106, 108])]


def test_read_permutations_without_header(make_xlsx):
    path = make_xlsx([[3, 1], [2, 2]])
    assert read_permutations(path) == [("Permuta_1", [3, 2]), ("Permuta_2", [1, 2])]


def test_read_permutations_skips_empty_columns(make_xlsx):
    path = make_xlsx([["A", None, "C"], [1, None, 2]])
    assert read_permutations(path) == [("A", [1]), ("C", [2])]


def test_read_permutations_rejects_text(make_xlsx):
    path = make_xlsx([["Permuta_1"], [5], ["five"]])
    with pytest.raises(ValueError, match="A3"):
        read_permutations(path)


# ----------------------------
# Reading the question bank
# ----------------------------

def test_bank_finds_all_questions(bank_path):
    bank = QuestionBank(str(bank_path))
    assert sorted(bank.questions) == [1, 2, 3, 4, 5, 6]
    assert bank.duplicates == []


def test_section_titles_are_never_copied(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [1, 2, 3, 4, 5, 6])
    assert not any(t.startswith("*") or "SECTION" in t or "SEMINARS" in t for t in texts(doc))


# ----------------------------
# Building exams
# ----------------------------

def test_questions_follow_requested_order_and_are_renumbered(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [4, 1, 3])
    assert question_lines(doc) == ["1. Fourth question", "2. First question / Primera pregunta", "3. Third question"]


def test_title_is_first(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [1])
    assert doc.paragraphs[0].text == "Test Exam"
    assert doc.paragraphs[0].style.name == "Title"


def test_number_split_over_runs_is_replaced(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [3, 2])
    assert "2. Second question with picture" in texts(doc)


def test_space_after_number_survives_in_word(bank_path, tmp_path):
    # Word ignores leading spaces unless xml:space="preserve" is set on the text node
    doc, _ = build(bank_path, tmp_path, [2])
    q = next(p for p in doc.paragraphs if p.text.startswith("1."))
    for t in q._p.iter(qn('w:t')):
        if t.text and t.text[0].isspace():
            assert t.get(qn('xml:space')) == "preserve"


def test_parenthesis_numbering_is_kept(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [1, 5])
    assert any(t.startswith("2) Fifth question") for t in texts(doc))


def test_source_formatting_is_kept(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [2, 1])
    q = next(p for p in doc.paragraphs if p.text.startswith("1. Second"))
    assert q.runs[-1].italic
    assert red_runs(doc) == ["Second right", "Right one"]
    assert num_ids(doc, "Second right")  # still an auto-lettered list


def test_answers_and_typed_lines_stay_with_their_question(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [5, 1])
    t = texts(doc)
    assert t[1].startswith("1) Fifth question") and "a. Fifth answer A" in t[1]
    assert t[2] == "e. Fifth answer E"
    assert t[3:] == ["2. First question / Primera pregunta", "Wrong one", "Right one"]


def test_each_question_restarts_its_lettering(bank_path, tmp_path):
    # Questions 3 and 4 share one list in the bank; in the exam they must not
    doc, _ = build(bank_path, tmp_path, [3, 4])
    q3, q4 = num_ids(doc, "Third"), num_ids(doc, "Fourth")
    assert len(set(q3)) == 1 and len(set(q4)) == 1 and q3[0] != q4[0]
    numbering = doc.part.numbering_part.element
    for num_id in (q3[0], q4[0]):
        num = numbering.num_having_numId(num_id)
        assert num.xpath('./w:lvlOverride[@w:ilvl="0"]/w:startOverride/@w:val') == ["1"]


def test_images_kept_and_unused_images_dropped(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [2])
    images = [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]
    assert len(images) == 1
    assert images[0].rId in {e.get(qn('r:embed')) for e in doc.element.body.iter(qn('a:blip'))}


def test_no_images_when_none_selected(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [1])
    assert not [r for r in doc.part.rels.values() if r.reltype == RT.IMAGE]


def test_remove_answer_marks(bank_path, tmp_path):
    doc, _ = build(bank_path, tmp_path, [1, 2], remove_answer_marks=True)
    assert red_runs(doc) == []
    assert "Right one" in texts(doc)


def test_report_missing_repeated_and_empty(bank_path, tmp_path):
    doc, report = build(bank_path, tmp_path, [1, 99, 1, 6])
    assert report == {"missing": [99], "repeated": [1], "empty": [6]}
    assert question_lines(doc) == ["1. First question / Primera pregunta", "2. "]
    assert "99" in describe_report(report) and "6" in describe_report(report)


def test_progress_called_per_number(bank_path, tmp_path):
    steps = []
    bank = QuestionBank(str(bank_path))
    build_exam(bank, [1, 99, 2], "T", str(tmp_path / "e.docx"), on_progress=steps.append)
    assert steps == [1, 2, 3]


def test_safe_filename():
    assert safe_filename('Exam: 1/2 "final"?') == "Exam_ 1_2 _final_"
    assert safe_filename("...") == "Exam"


# ----------------------------
# Real question bank (optional)
# ----------------------------

POOL = os.environ.get("SYNAPSE_POOL")
PERMUTAS = os.environ.get("SYNAPSE_PERMUTAS")


@pytest.mark.skipif(not (POOL and PERMUTAS), reason="set SYNAPSE_POOL and SYNAPSE_PERMUTAS to run against the real files")
def test_real_permutations(tmp_path):
    bank = QuestionBank(POOL)
    source = Document(POOL).element.body
    children = list(source.iterchildren())
    def stem(n):
        return re.sub(r"^\s*\d+\s*[\.\)]", "", "".join(t.text or "" for t in children[bank.questions[n][0]].iter(qn('w:t')))).strip()

    for name, numbers in read_permutations(PERMUTAS):
        out = tmp_path / f"{name}.docx"
        report = build_exam(bank, numbers, name, str(out))
        assert not describe_report(report), describe_report(report)
        lines = question_lines(Document(str(out)))
        assert len(lines) == len(numbers)
        for i, (line, n) in enumerate(zip(lines, numbers), start=1):
            assert line.startswith(str(i)) and line.split(None, 1)[-1].strip().startswith(stem(n)[:30].strip())
        assert not any(t.startswith("*") for t in texts(Document(str(out))))
